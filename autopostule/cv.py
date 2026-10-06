"""Lecture du CV (PDF, DOCX, TXT, MD) et extraction des informations utiles."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from .referentiel import METIERS, normaliser, toutes_les_competences

RE_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
RE_TELEPHONE = re.compile(r"(?:\+33\s?|0)[1-9](?:[\s.-]?\d{2}){4}")
RE_LINKEDIN = re.compile(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[\w%-]+/?", re.I)
RE_GITHUB = re.compile(r"(?:https?://)?github\.com/[\w-]+/?", re.I)


@dataclass
class AnalyseCV:
    texte: str
    email: str | None = None
    telephone: str | None = None
    linkedin: str | None = None
    github: str | None = None
    competences: list[str] = field(default_factory=list)
    scores_metiers: dict[str, int] = field(default_factory=dict)

    @property
    def metiers_suggeres(self) -> list[str]:
        return [m for m, s in sorted(self.scores_metiers.items(), key=lambda x: -x[1]) if s > 0]

    def competences_pour(self, cle_metier: str, maximum: int = 6) -> list[str]:
        """Compétences du CV pertinentes pour un métier, dans l'ordre du référentiel."""
        metier = METIERS[cle_metier]
        dans_cv = {normaliser(c) for c in self.competences}
        return [c for c in metier.competences if normaliser(c) in dans_cv][:maximum]


def extraire_texte(chemin: Path) -> str:
    chemin = Path(chemin)
    if not chemin.exists():
        raise FileNotFoundError(f"CV introuvable : {chemin}")
    suffixe = chemin.suffix.lower()
    if suffixe == ".pdf":
        from pypdf import PdfReader

        lecteur = PdfReader(str(chemin))
        return "\n".join((page.extract_text() or "") for page in lecteur.pages)
    if suffixe == ".docx":
        import docx

        document = docx.Document(str(chemin))
        lignes = [p.text for p in document.paragraphs]
        for table in document.tables:
            for ligne in table.rows:
                lignes.append(" | ".join(cell.text for cell in ligne.cells))
        return "\n".join(lignes)
    if suffixe in {".txt", ".md", ".markdown"}:
        return chemin.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Format de CV non pris en charge : {suffixe} (PDF, DOCX, TXT ou MD)")


def _contient(texte_normalise: str, competence: str) -> bool:
    terme = normaliser(competence)
    # Bornes de mot souples : « ci/cd », « .net », « tcp/ip » contiennent des caractères non alphanumériques.
    motif = r"(?<![\w])" + re.escape(terme) + r"(?![\w])"
    return re.search(motif, texte_normalise) is not None


def detecter_competences(texte: str) -> list[str]:
    normalise = normaliser(texte)
    return [c for c in toutes_les_competences() if _contient(normalise, c)]


def analyser(texte: str) -> AnalyseCV:
    analyse = AnalyseCV(texte=texte)
    if m := RE_EMAIL.search(texte):
        analyse.email = m.group(0)
    if m := RE_TELEPHONE.search(texte):
        analyse.telephone = m.group(0)
    if m := RE_LINKEDIN.search(texte):
        analyse.linkedin = m.group(0)
    if m := RE_GITHUB.search(texte):
        analyse.github = m.group(0)
    analyse.competences = detecter_competences(texte)
    trouvees = {normaliser(c) for c in analyse.competences}
    for cle, metier in METIERS.items():
        analyse.scores_metiers[cle] = sum(1 for c in metier.competences if normaliser(c) in trouvees)
    return analyse


def analyser_fichier(chemin: Path) -> AnalyseCV:
    return analyser(extraire_texte(chemin))
